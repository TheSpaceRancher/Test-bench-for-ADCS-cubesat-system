from __future__ import annotations

import argparse
import queue
import threading
from datetime import datetime
from pathlib import Path
from typing import Dict, Optional

import numpy as np
import matplotlib.pyplot as plt
from matplotlib.animation import FuncAnimation
from matplotlib.widgets import Button
from mpl_toolkits.mplot3d.art3d import Poly3DCollection
import serial
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter


FIELDS = [
    "t_ms", "Bx_uT", "By_uT", "Bz_uT",
    "dBx_uT_s", "dBy_uT_s", "dBz_uT_s",
    "I_X_est_prom_A", "I_Y_est_prom_A", "duty_X", "duty_Y", "seq",
]
MAX_L298N_CURRENT_A = 2.0

def parse_data_line(line: str) -> Optional[Dict[str, float]]:
    if not line.startswith("DATOS:"):
        return None
    values: Dict[str, float] = {}
    for token in line[len("DATOS:"):].strip().split(","):
        if "=" not in token:
            continue
        key, raw_value = token.split("=", 1)
        try:
            values[key.strip()] = float(raw_value.strip())
        except ValueError:
            return None
    if not all(field in values for field in FIELDS):
        return None
    values["I_X_est_prom_A"] = min(
        max(values["I_X_est_prom_A"], 0.0), MAX_L298N_CURRENT_A
    )
    values["I_Y_est_prom_A"] = min(
        max(values["I_Y_est_prom_A"], 0.0), MAX_L298N_CURRENT_A
    )
    return values


def serial_reader(
    serial_port: str,
    baud_rate: int,
    samples: "queue.Queue[Dict[str, float]]",
    stop_event: threading.Event,
) -> None:
    try:
        with serial.Serial(serial_port, baud_rate, timeout=0.5) as connection:
            while not stop_event.is_set():
                line = connection.readline().decode("utf-8", errors="replace").strip()
                if line:
                    sample = parse_data_line(line)
                    if sample is not None:
                        samples.put(sample)
    except (serial.SerialException, OSError) as error:
        samples.put({"_reader_error": str(error)})


def create_workbook(path: Path) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Telemetry"
    sheet.append(FIELDS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(FIELDS))}1"
    for cell in sheet[1]:
        cell.font = cell.font.copy(bold=True)
    workbook.save(path)
    return workbook


def add_excel_charts(workbook: Workbook) -> None:
    sheet = workbook["Telemetry"]
    last_row = sheet.max_row
    if last_row < 2:
        return
    columns = {cell.value: cell.column for cell in sheet[1]}
    chart_specs = [
        ("Campo magnetico", ["Bx_uT", "By_uT", "Bz_uT"], "uT"),
        ("Derivada del campo", ["dBx_uT_s", "dBy_uT_s", "dBz_uT_s"], "uT/s"),
        ("Corriente estimada", ["I_X_est_prom_A", "I_Y_est_prom_A"], "A"),
    ]
    for offset, (title, names, y_title) in enumerate(chart_specs):
        chart = LineChart()
        chart.title = title
        chart.y_axis.title = y_title
        chart.x_axis.title = "Tiempo (ms)"
        chart.height = 7
        chart.width = 14
        for name in names:
            data = Reference(sheet, min_col=columns[name], min_row=1, max_row=last_row)
            chart.add_data(data, titles_from_data=True)
        categories = Reference(
            sheet, min_col=columns["t_ms"], min_row=2, max_row=last_row
        )
        chart.set_categories(categories)
        sheet.add_chart(chart, f"A{last_row + 3 + offset * 16}")


def rotation_from_vectors(source: np.ndarray, target: np.ndarray) -> np.ndarray:
    """Devuelve la rotacion minima que lleva source hasta target."""
    source = source / np.linalg.norm(source)
    target = target / np.linalg.norm(target)
    cross = np.cross(source, target)
    dot = float(np.clip(np.dot(source, target), -1.0, 1.0))
    if np.linalg.norm(cross) < 1e-9:
        if dot > 0:
            return np.eye(3)
        perpendicular = np.array([1.0, 0.0, 0.0])
        if abs(source[0]) > 0.9:
            perpendicular = np.array([0.0, 1.0, 0.0])
        axis = perpendicular - source * np.dot(perpendicular, source)
        axis /= np.linalg.norm(axis)
        return 2.0 * np.outer(axis, axis) - np.eye(3)
    skew = np.array([
        [0.0, -cross[2], cross[1]],
        [cross[2], 0.0, -cross[0]],
        [-cross[1], cross[0], 0.0],
    ])
    return np.eye(3) + skew + skew @ skew * ((1.0 - dot) / np.dot(cross, cross))


def draw_satellite(scene, rotation: np.ndarray, field_vector: np.ndarray) -> None:
    legend = scene.get_legend()
    if legend is not None:
        legend.remove()
    for artist in list(scene.collections) + list(scene.lines) + list(scene.texts):
        artist.remove()
    vertices = np.array([
        [-1, -1, -1], [1, -1, -1], [1, 1, -1], [-1, 1, -1],
        [-1, -1, 1], [1, -1, 1], [1, 1, 1], [-1, 1, 1],
    ], dtype=float) * 0.8
    vertices = vertices @ rotation.T
    faces = [
        [0, 1, 2, 3], [4, 5, 6, 7], [0, 1, 5, 4],
        [2, 3, 7, 6], [1, 2, 6, 5], [0, 3, 7, 4],
    ]
    scene.add_collection3d(Poly3DCollection(
        [[vertices[index] for index in face] for face in faces],
        alpha=0.25, facecolor="steelblue", edgecolor="black",
    ))
    axis_colors = ("red", "green", "blue")
    axis_names = ("X", "Y", "Z")
    for axis, color, name in zip(np.eye(3), axis_colors, axis_names):
        endpoint = rotation @ (axis * 1.5)
        scene.plot(
            [0, endpoint[0]], [0, endpoint[1]], [0, endpoint[2]],
            color=color, linewidth=2.5, label=f"Eje {name}",
        )
        scene.text(endpoint[0], endpoint[1], endpoint[2], name, color=color)
    scene.quiver(
        0, 0, 0, field_vector[0], field_vector[1], field_vector[2],
        color="darkorange", linewidth=3, arrow_length_ratio=0.15,
    )
    scene.text(
        field_vector[0], field_vector[1], field_vector[2],
        " B ", color="darkorange",
    )
    scene.legend(loc="upper left", fontsize=8)


def main() -> None:
    parser = argparse.ArgumentParser(description="Dashboard de telemetria ESP-NOW")
    parser.add_argument("--port", required=True, help="Puerto del ESP32 receptor, por ejemplo COM5")
    parser.add_argument("--baud", type=int, default=115200)
    parser.add_argument("--output", default="telemetria_bdot.xlsx")
    parser.add_argument("--window", type=int, default=600)
    args = parser.parse_args()

    samples: "queue.Queue[Dict[str, float]]" = queue.Queue()
    stop_event = threading.Event()
    reader = threading.Thread(
        target=serial_reader,
        args=(args.port, args.baud, samples, stop_event),
        daemon=True,
    )
    reader.start()

    history = {field: [] for field in FIELDS}
    recording = False
    workbook: Optional[Workbook] = None
    sheet = None
    output_path: Optional[Path] = None
    saved_count = 0

    figure = plt.figure(figsize=(13, 8))
    figure.subplots_adjust(bottom=0.14, left=0.03, right=0.98, wspace=0.12)
    grid = figure.add_gridspec(1, 2, width_ratios=(3.0, 1.0))
    scene = figure.add_subplot(grid[0, 0], projection="3d")
    numeric_panel = figure.add_subplot(grid[0, 1])
    try:
        figure.canvas.manager.set_window_title("Dashboard B-dot - ESP-NOW")
    except AttributeError:
        pass

    numeric_panel.axis("off")
    # Referencia visual: el campo terrestre apunta horizontalmente hacia +X.
    earth_field = np.array([1.0, 0.0, 0.0])
    scene.set_title("Actitud estimada")
    scene.set_xlim(-2.0, 2.0)
    scene.set_ylim(-2.0, 2.0)
    scene.set_zlim(-2.0, 2.0)
    scene.set_box_aspect((1, 1, 1))
    scene.set_xlabel("Marco terrestre X")
    scene.set_ylabel("Marco terrestre Y")
    scene.set_zlabel("Marco terrestre Z")

    status = figure.text(0.5, 0.02, "GRABACION DETENIDA", ha="center")
    start_axis = figure.add_axes((0.25, 0.045, 0.20, 0.045))
    stop_axis = figure.add_axes((0.55, 0.045, 0.20, 0.045))
    start_button = Button(start_axis, "Comenzar a grabar")
    stop_button = Button(stop_axis, "Parar grabacion")

    def start_recording(_event) -> None:
        nonlocal recording, workbook, sheet, output_path, saved_count
        if recording:
            return
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        requested = Path(args.output)
        output_path = requested.with_name(
            f"{requested.stem}_{timestamp}{requested.suffix or '.xlsx'}"
        ).resolve()
        workbook = create_workbook(output_path)
        sheet = workbook["Telemetry"]
        saved_count = 0
        recording = True
        status.set_text(f"GRABANDO: {output_path.name}")
        status.set_color("green")

    def stop_recording(_event) -> None:
        nonlocal recording
        if not recording or workbook is None or output_path is None:
            return
        recording = False
        add_excel_charts(workbook)
        workbook.save(output_path)
        status.set_text(f"GRABACION DETENIDA | Guardado: {output_path.name}")
        status.set_color("black")

    start_button.on_clicked(start_recording)
    stop_button.on_clicked(stop_recording)

    def update(_frame: int):
        nonlocal saved_count
        try:
            while True:
                sample = samples.get_nowait()
                if "_reader_error" in sample:
                    status.set_text(f"ERROR SERIAL: {sample['_reader_error']}")
                    status.set_color("red")
                    continue
                for field in FIELDS:
                    history[field].append(sample[field])
                if recording and sheet is not None and workbook is not None and output_path:
                    sheet.append([sample[field] for field in FIELDS])
                    saved_count += 1
                    if saved_count % 10 == 0:
                        workbook.save(output_path)
        except queue.Empty:
            pass

        if history["Bx_uT"]:
            received_field = np.array([
                history["Bx_uT"][-1],
                history["By_uT"][-1],
                history["Bz_uT"][-1],
            ])
            measured_field = received_field
            displayed_derivative = np.array([
                history["dBx_uT_s"][-1],
                history["dBy_uT_s"][-1],
                history["dBz_uT_s"][-1],
            ])
            magnitude = np.linalg.norm(measured_field)
            if magnitude > 1e-6:
                rotation = rotation_from_vectors(measured_field, earth_field)
                draw_satellite(scene, rotation, earth_field)
                numeric_panel.clear()
                numeric_panel.axis("off")
                received_sample = {
                    field: history[field][-1]
                    for field in (
                        "t_ms", "Bx_uT", "By_uT", "Bz_uT",
                        "dBx_uT_s", "dBy_uT_s", "dBz_uT_s",
                        "I_X_est_prom_A", "I_Y_est_prom_A",
                    )
                }
                numeric_panel.text(
                    0.0, 0.98,
                    "DATOS DEL SATELITE\n\n"
                    f"Tiempo: {received_sample['t_ms']:.0f} ms\n\n"
                    "MAGNETOMETRO\n"
                    f"Bx: {measured_field[0]:.3f} uT\n"
                    f"By: {measured_field[1]:.3f} uT\n"
                    f"Bz: {measured_field[2]:.3f} uT\n"
                    f"|B|: {magnitude:.3f} uT\n\n"
                    "DERIVADA dB/dt\n"
                    f"dBx/dt: {displayed_derivative[0]:.3f} uT/s\n"
                    f"dBy/dt: {displayed_derivative[1]:.3f} uT/s\n"
                    f"dBz/dt: {displayed_derivative[2]:.3f} uT/s\n\n"
                    "CORRIENTE ESTIMADA\n"
                    f"MT-X: {received_sample['I_X_est_prom_A']:.3f} A\n"
                    f"MT-Y: {received_sample['I_Y_est_prom_A']:.3f} A\n\n"
                    "B fijo: naranja\n"
                    "X rojo | Y verde | Z azul\n",
                    va="top", family="monospace", fontsize=9,
                )
        return ()

    animation = FuncAnimation(
        figure, update, interval=100, blit=False, cache_frame_data=False
    )
    try:
        plt.show()
    finally:
        if recording and workbook is not None and output_path is not None:
            add_excel_charts(workbook)
            workbook.save(output_path)
        animation.event_source.stop()
        stop_event.set()
        reader.join(timeout=2)


if __name__ == "__main__":
    main()
